#!/usr/bin/env python3
"""Generate Kelvin Oladotun Esanoluwa CV as DOCX and PDF."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, mm
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

OUT_DIR = Path(__file__).resolve().parent
NAVY = RGBColor(0x1B, 0x36, 0x5D)
NAVY_HEX = colors.HexColor("#1B365D")
SLATE = colors.HexColor("#334155")
RULE = colors.HexColor("#1B365D")
LIGHT = colors.HexColor("#64748B")


def set_run_font(run, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_bottom_border(paragraph, color="1B365D", size="12"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def tight_paragraph(doc, space_after=4, space_before=0, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    return p


def add_section(doc, title):
    p = tight_paragraph(doc, space_after=6, space_before=10)
    run = p.add_run(title.upper())
    set_run_font(run, size=12, bold=True, color=NAVY)
    add_bottom_border(p)
    return p


def add_bullet(doc, text):
    p = tight_paragraph(doc, space_after=3)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    run = p.add_run("•  " + text)
    set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))
    return p


def build_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    name = tight_paragraph(doc, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = name.add_run("KELVIN OLADOTUN ESANOLUWA")
    set_run_font(run, size=20, bold=True, color=NAVY)

    title = tight_paragraph(doc, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = title.add_run(
        "Budget Analyst  ·  Financial Planning  ·  Quantitative Risk  ·  Actuarial Science"
    )
    set_run_font(run, size=11, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    contact = tight_paragraph(doc, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = contact.add_run(
        "17 Ugochukwu Orji Street, Westgate Estate, Igbo Efon, Lagos, Nigeria"
    )
    set_run_font(run, size=10, color=RGBColor(0x33, 0x41, 0x55))

    contact2 = tight_paragraph(doc, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = contact2.add_run(
        "+234 812 403 1683   ·   Dotun_esanoluwa@yahoo.com   ·   Open to full-time roles in Lagos"
    )
    set_run_font(run, size=10, color=RGBColor(0x33, 0x41, 0x55))

    add_section(doc, "Professional Summary")
    summary = tight_paragraph(doc, space_after=4)
    run = summary.add_run(
        "Industrial Mathematics graduate and actuarial science postgraduate with hands-on "
        "public-sector experience in budget formulation, expenditure control, cost analysis, "
        "and stakeholder reporting. Combines a strong numerical foundation with practical "
        "work at the Lagos State House of Assembly and its Service Commission, plus early "
        "commercial exposure in airline procurement. Comfortable moving between "
        "spreadsheets, financial schedules, and clear written advice. Seeking a Budget "
        "Analyst, Financial Analyst, Risk Analyst, or Actuarial Analyst role where careful "
        "quantification and reliable documentation support better funding and risk decisions."
    )
    set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))

    add_section(doc, "Core Qualifications")
    quals = [
        "Budget formulation, allocation scheduling, and in-year expenditure monitoring for designated public votes.",
        "Cost, variance, and management-information analysis that turns financial reports into decision-ready briefings.",
        "Quantitative methods from industrial mathematics and actuarial postgraduate study: probability, forecasting, financial mathematics, and risk assessment.",
        "Document control and data validation — checking source figures, reconciling returns, and keeping audit-ready working papers.",
        "Stakeholder communication with supervisors, directorates, and non-finance users who need plain-language budget clarification.",
        "Applied Excel modelling for projections, lookups, pivot summaries, charts, and purchase or budget tracking workbooks.",
        "Research discipline across qualitative and quantitative sources, with accurate written findings and formal cost notes.",
        "Procurement support skills: quotations, cost parameters, vendor records, and inventory-related purchase documentation.",
    ]
    for item in quals:
        add_bullet(doc, item)

    add_section(doc, "Professional Experience")

    role = tight_paragraph(doc, space_after=0)
    run = role.add_run("Budget Officer")
    set_run_font(run, size=11.5, bold=True, color=NAVY)
    run = role.add_run("  —  Lagos State House of Assembly")
    set_run_font(run, size=11, color=RGBColor(0x1E, 0x29, 0x3B))

    meta = tight_paragraph(doc, space_after=4)
    run = meta.add_run("May 2017 – June 2018  ·  Lagos, Nigeria  ·  Full-time")
    set_run_font(run, size=10, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    for item in [
        "Supported the annual budgeting and financial planning cycle by preparing allocation schedules and expense projections for designated votes.",
        "Interpreted appropriation and management reports so supervisors had timely, accurate figures for funding and reallocation decisions.",
        "Monitored balances and expenditure against approved provisions, helping protect the financial integrity of assigned budget lines.",
        "Coordinated budgetary and statistical data from contributing units, reconciled inconsistencies, and maintained clear working papers.",
        "Produced cost analyses of recurrent and capital spend and issued formal written findings for internal review.",
        "Briefed stakeholders on budget status, implementation questions, and the practical meaning of figures in management reports.",
    ]:
        add_bullet(doc, item)

    role = tight_paragraph(doc, space_after=0, space_before=8)
    run = role.add_run("Budget Officer (National Youth Service)")
    set_run_font(run, size=11.5, bold=True, color=NAVY)
    run = role.add_run("  —  Lagos State House of Assembly Service Commission")
    set_run_font(run, size=11, color=RGBColor(0x1E, 0x29, 0x3B))

    meta = tight_paragraph(doc, space_after=4)
    run = meta.add_run("September 2014 – September 2015  ·  Lagos, Nigeria  ·  NYSC")
    set_run_font(run, size=10, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    for item in [
        "Prepared financial schedules and supporting documents used in state budget compilation and internal review.",
        "Checked the validity, completeness, and consistency of figures supplied in financial documents before onward use.",
        "Worked with other departments to clarify information needed for effective implementation of budgets and budgetary policy.",
        "Built an early working knowledge of public-sector budget language, vote control, and inter-departmental reporting.",
    ]:
        add_bullet(doc, item)

    role = tight_paragraph(doc, space_after=0, space_before=8)
    run = role.add_run("Procurement Officer (Industrial Internship)")
    set_run_font(run, size=11.5, bold=True, color=NAVY)
    run = role.add_run("  —  Arik Air Plc")
    set_run_font(run, size=11, color=RGBColor(0x1E, 0x29, 0x3B))

    meta = tight_paragraph(doc, space_after=4)
    run = meta.add_run("March 2009 – August 2009  ·  Lagos, Nigeria  ·  Internship")
    set_run_font(run, size=10, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    for item in [
        "Improved day-to-day retrieval of procurement records by sorting documentation into a clearer, more accessible filing system.",
        "Assisted with acquisition of operational items, including gathering quotations and supporting purchase documentation.",
        "Helped estimate and record cost parameters and purchase budgets so orders stayed within agreed limits.",
        "Built spreadsheets to track supplier information, order status, and relevant commercial details for the team.",
        "Supported price and fee comparisons while helping the function maintain continuity of required services.",
    ]:
        add_bullet(doc, item)

    add_section(doc, "Education")

    edu = tight_paragraph(doc, space_after=0)
    run = edu.add_run("M.Sc. Actuarial Science (Part-time)")
    set_run_font(run, size=11.5, bold=True, color=NAVY)
    run = edu.add_run("  —  University of Lagos")
    set_run_font(run, size=11, color=RGBColor(0x1E, 0x29, 0x3B))
    meta = tight_paragraph(doc, space_after=2)
    run = meta.add_run("Commenced 2020  ·  In progress")
    set_run_font(run, size=10, italic=True, color=RGBColor(0x47, 0x55, 0x69))
    note = tight_paragraph(doc, space_after=6)
    run = note.add_run(
        "Postgraduate study covering probability and statistical inference, financial mathematics, "
        "life contingencies, survival models, risk theory, insurance mathematics, and quantitative "
        "modelling for insurance and financial risk. Builds directly on undergraduate mathematics "
        "and on practical budget and cost-analysis work."
    )
    set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))

    edu = tight_paragraph(doc, space_after=0)
    run = edu.add_run("B.Sc. Industrial Mathematics  —  Second Class Honours")
    set_run_font(run, size=11.5, bold=True, color=NAVY)
    run = edu.add_run("  —  Covenant University, Ota, Ogun State")
    set_run_font(run, size=11, color=RGBColor(0x1E, 0x29, 0x3B))
    meta = tight_paragraph(doc, space_after=2)
    run = meta.add_run("Awarded 2014")
    set_run_font(run, size=10, italic=True, color=RGBColor(0x47, 0x55, 0x69))
    note = tight_paragraph(doc, space_after=6)
    run = note.add_run(
        "Core training in calculus, linear algebra, differential equations, probability and "
        "statistics, operations research, numerical analysis, mathematical modelling, and "
        "financial mathematics. Developed the quantitative habits used later in budgeting, "
        "forecasting, and risk analysis."
    )
    set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))

    edu = tight_paragraph(doc, space_after=0)
    run = edu.add_run("West African Senior School Certificate")
    set_run_font(run, size=11.5, bold=True, color=NAVY)
    run = edu.add_run("  —  Dansol High School, Agidingbi, Ikeja")
    set_run_font(run, size=11, color=RGBColor(0x1E, 0x29, 0x3B))
    meta = tight_paragraph(doc, space_after=6)
    run = meta.add_run(
        "WASSCE: 2 Distinctions, 3 Credits, 4 Passes  ·  Junior WAEC: 5 Distinctions, 6 Credits"
    )
    set_run_font(run, size=10, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    add_section(doc, "Technical Skills")
    skills = [
        ("Budgeting & finance", "Budget preparation, allocation control, expenditure monitoring, cost analysis, financial planning, forecasting, variance review, management reporting."),
        ("Quantitative methods", "Financial mathematics, probability, statistical description, quantitative risk assessment, business research, qualitative and quantitative analysis."),
        ("Tools", "Microsoft Excel (lookups, pivot tables, charts, financial schedules), Word, PowerPoint; coursework exposure to statistical packages used in actuarial study (SPSS / R)."),
        ("Operations", "Documentation control, inventory-related cost parameters, procurement records, data validation, formal report writing."),
        ("Working style", "Accuracy, attention to detail, independent thinking, stakeholder briefing, and a steady habit of checking figures before they are issued."),
    ]
    for label, text in skills:
        p = tight_paragraph(doc, space_after=3)
        run = p.add_run(label + ": ")
        set_run_font(run, size=10.5, bold=True, color=NAVY)
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))

    add_section(doc, "Additional Information")
    extras = [
        "Language: English — professional written and spoken communication for reports, briefings, and correspondence.",
        "National Youth Service completed (2014–2015) with a budget posting at the Lagos State House of Assembly Service Commission.",
        "Applied knowledge areas: time value of money and discounted cash-flow reasoning; probability and descriptive statistics; public appropriation and vote-head monitoring; cost estimation for operational purchases.",
        "Work authorisation: Nigerian citizen, based in Lagos and available for full-time office or hybrid roles.",
    ]
    for item in extras:
        add_bullet(doc, item)

    add_section(doc, "Interests")
    p = tight_paragraph(doc, space_after=4)
    run = p.add_run(
        "Financial magazines and educative articles; philosophy and motivational reading; "
        "novels; writing poetry. Continued personal study in risk, insurance mathematics, "
        "and practical financial analysis."
    )
    set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))

    add_section(doc, "References")
    p = tight_paragraph(doc, space_after=0)
    run = p.add_run("Available on request.")
    set_run_font(run, size=10.5, color=RGBColor(0x1E, 0x29, 0x3B))

    doc.save(path)


def heading_style(name, parent, **kwargs):
    style = ParagraphStyle(name, parent=parent, **kwargs)
    return style


def build_pdf(path: Path) -> None:
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=1.55 * cm,
        rightMargin=1.55 * cm,
        topMargin=1.35 * cm,
        bottomMargin=1.35 * cm,
        title="Kelvin Oladotun Esanoluwa — Curriculum Vitae",
        author="Kelvin Oladotun Esanoluwa",
    )
    styles = getSampleStyleSheet()
    name = heading_style(
        "Name",
        styles["Normal"],
        fontName="Times-Bold",
        fontSize=18,
        leading=22,
        textColor=NAVY_HEX,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    role = heading_style(
        "RoleLine",
        styles["Normal"],
        fontName="Times-Italic",
        fontSize=10,
        leading=13,
        textColor=LIGHT,
        alignment=TA_CENTER,
        spaceAfter=3,
    )
    contact = heading_style(
        "Contact",
        styles["Normal"],
        fontName="Times-Roman",
        fontSize=9.2,
        leading=12,
        textColor=SLATE,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    section = heading_style(
        "Section",
        styles["Normal"],
        fontName="Times-Bold",
        fontSize=11,
        leading=14,
        textColor=NAVY_HEX,
        spaceBefore=10,
        spaceAfter=3,
        tracking=0.4,
    )
    body = heading_style(
        "Body",
        styles["Normal"],
        fontName="Times-Roman",
        fontSize=9.6,
        leading=12.6,
        textColor=SLATE,
        alignment=TA_JUSTIFY,
        spaceAfter=3,
    )
    job_title = heading_style(
        "JobTitle",
        styles["Normal"],
        fontName="Times-Bold",
        fontSize=10.4,
        leading=13,
        textColor=NAVY_HEX,
        spaceBefore=7,
        spaceAfter=0,
    )
    job_meta = heading_style(
        "JobMeta",
        styles["Normal"],
        fontName="Times-Italic",
        fontSize=9,
        leading=12,
        textColor=LIGHT,
        spaceAfter=3,
    )
    bullet = heading_style(
        "Bullet",
        styles["Normal"],
        fontName="Times-Roman",
        fontSize=9.6,
        leading=12.4,
        textColor=SLATE,
        leftIndent=12,
        firstLineIndent=-10,
        spaceAfter=2,
    )
    skill_label = heading_style(
        "SkillLabel",
        styles["Normal"],
        fontName="Times-Bold",
        fontSize=9.6,
        leading=12.6,
        textColor=NAVY_HEX,
    )

    story = []
    story.append(Paragraph("KELVIN OLADOTUN ESANOLUWA", name))
    story.append(
        Paragraph(
            "Budget Analyst  ·  Financial Planning  ·  Quantitative Risk  ·  Actuarial Science",
            role,
        )
    )
    story.append(
        Paragraph(
            "17 Ugochukwu Orji Street, Westgate Estate, Igbo Efon, Lagos, Nigeria",
            contact,
        )
    )
    story.append(
        Paragraph(
            "+234 812 403 1683&nbsp;&nbsp;·&nbsp;&nbsp;Dotun_esanoluwa@yahoo.com&nbsp;&nbsp;·&nbsp;&nbsp;Open to full-time roles in Lagos",
            contact,
        )
    )
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1.2, color=RULE, spaceAfter=2))

    def add_sec(title):
        story.append(Paragraph(title.upper(), section))
        story.append(HRFlowable(width="100%", thickness=0.5, color=RULE, spaceAfter=6))

    add_sec("Professional Summary")
    story.append(
        Paragraph(
            "Industrial Mathematics graduate and actuarial science postgraduate with hands-on "
            "public-sector experience in budget formulation, expenditure control, cost analysis, "
            "and stakeholder reporting. Combines a strong numerical foundation with practical "
            "work at the Lagos State House of Assembly and its Service Commission, plus early "
            "commercial exposure in airline procurement. Comfortable moving between spreadsheets, "
            "financial schedules, and clear written advice. Seeking a Budget Analyst, Financial "
            "Analyst, Risk Analyst, or Actuarial Analyst role where careful quantification and "
            "reliable documentation support better funding and risk decisions.",
            body,
        )
    )

    add_sec("Core Qualifications")
    for item in [
        "Budget formulation, allocation scheduling, and in-year expenditure monitoring for designated public votes.",
        "Cost, variance, and management-information analysis that turns financial reports into decision-ready briefings.",
        "Quantitative methods from industrial mathematics and actuarial postgraduate study: probability, forecasting, financial mathematics, and risk assessment.",
        "Document control and data validation — checking source figures, reconciling returns, and keeping audit-ready working papers.",
        "Stakeholder communication with supervisors, directorates, and non-finance users who need plain-language budget clarification.",
        "Applied Excel modelling for projections, lookups, pivot summaries, charts, and purchase or budget tracking workbooks.",
        "Research discipline across qualitative and quantitative sources, with accurate written findings and formal cost notes.",
        "Procurement support skills: quotations, cost parameters, vendor records, and inventory-related purchase documentation.",
    ]:
        story.append(Paragraph("•  " + item, bullet))

    add_sec("Professional Experience")
    story.append(Paragraph("Budget Officer  —  Lagos State House of Assembly", job_title))
    story.append(Paragraph("May 2017 – June 2018  ·  Lagos, Nigeria  ·  Full-time", job_meta))
    for item in [
        "Supported the annual budgeting and financial planning cycle by preparing allocation schedules and expense projections for designated votes.",
        "Interpreted appropriation and management reports so supervisors had timely, accurate figures for funding and reallocation decisions.",
        "Monitored balances and expenditure against approved provisions, helping protect the financial integrity of assigned budget lines.",
        "Coordinated budgetary and statistical data from contributing units, reconciled inconsistencies, and maintained clear working papers.",
        "Produced cost analyses of recurrent and capital spend and issued formal written findings for internal review.",
        "Briefed stakeholders on budget status, implementation questions, and the practical meaning of figures in management reports.",
    ]:
        story.append(Paragraph("•  " + item, bullet))

    story.append(
        Paragraph(
            "Budget Officer (National Youth Service)  —  Lagos State House of Assembly Service Commission",
            job_title,
        )
    )
    story.append(Paragraph("September 2014 – September 2015  ·  Lagos, Nigeria  ·  NYSC", job_meta))
    for item in [
        "Prepared financial schedules and supporting documents used in state budget compilation and internal review.",
        "Checked the validity, completeness, and consistency of figures supplied in financial documents before onward use.",
        "Worked with other departments to clarify information needed for effective implementation of budgets and budgetary policy.",
        "Built an early working knowledge of public-sector budget language, vote control, and inter-departmental reporting.",
    ]:
        story.append(Paragraph("•  " + item, bullet))

    story.append(Paragraph("Procurement Officer (Industrial Internship)  —  Arik Air Plc", job_title))
    story.append(Paragraph("March 2009 – August 2009  ·  Lagos, Nigeria  ·  Internship", job_meta))
    for item in [
        "Improved day-to-day retrieval of procurement records by sorting documentation into a clearer, more accessible filing system.",
        "Assisted with acquisition of operational items, including gathering quotations and supporting purchase documentation.",
        "Helped estimate and record cost parameters and purchase budgets so orders stayed within agreed limits.",
        "Built spreadsheets to track supplier information, order status, and relevant commercial details for the team.",
        "Supported price and fee comparisons while helping the function maintain continuity of required services.",
    ]:
        story.append(Paragraph("•  " + item, bullet))

    add_sec("Education")
    story.append(Paragraph("M.Sc. Actuarial Science (Part-time)  —  University of Lagos", job_title))
    story.append(Paragraph("Commenced 2020  ·  In progress", job_meta))
    story.append(
        Paragraph(
            "Postgraduate study covering probability and statistical inference, financial mathematics, "
            "life contingencies, survival models, risk theory, insurance mathematics, and quantitative "
            "modelling for insurance and financial risk. Builds directly on undergraduate mathematics "
            "and on practical budget and cost-analysis work.",
            body,
        )
    )
    story.append(
        Paragraph(
            "B.Sc. Industrial Mathematics — Second Class Honours  —  Covenant University, Ota, Ogun State",
            job_title,
        )
    )
    story.append(Paragraph("Awarded 2014", job_meta))
    story.append(
        Paragraph(
            "Core training in calculus, linear algebra, differential equations, probability and "
            "statistics, operations research, numerical analysis, mathematical modelling, and "
            "financial mathematics. Developed the quantitative habits used later in budgeting, "
            "forecasting, and risk analysis.",
            body,
        )
    )
    story.append(Paragraph("West African Senior School Certificate  —  Dansol High School, Agidingbi, Ikeja", job_title))
    story.append(
        Paragraph(
            "WASSCE: 2 Distinctions, 3 Credits, 4 Passes  ·  Junior WAEC: 5 Distinctions, 6 Credits",
            job_meta,
        )
    )

    add_sec("Technical Skills")
    skill_rows = [
        ("<b>Budgeting &amp; finance:</b> Budget preparation, allocation control, expenditure monitoring, cost analysis, financial planning, forecasting, variance review, management reporting."),
        ("<b>Quantitative methods:</b> Financial mathematics, probability, statistical description, quantitative risk assessment, business research, qualitative and quantitative analysis."),
        ("<b>Tools:</b> Microsoft Excel (lookups, pivot tables, charts, financial schedules), Word, PowerPoint; coursework exposure to statistical packages used in actuarial study (SPSS / R)."),
        ("<b>Operations:</b> Documentation control, inventory-related cost parameters, procurement records, data validation, formal report writing."),
        ("<b>Working style:</b> Accuracy, attention to detail, independent thinking, stakeholder briefing, and a steady habit of checking figures before they are issued."),
    ]
    for row in skill_rows:
        story.append(Paragraph(row, body))

    add_sec("Additional Information")
    for item in [
        "Language: English — professional written and spoken communication for reports, briefings, and correspondence.",
        "National Youth Service completed (2014–2015) with a budget posting at the Lagos State House of Assembly Service Commission.",
        "Applied knowledge areas: time value of money and discounted cash-flow reasoning; probability and descriptive statistics; public appropriation and vote-head monitoring; cost estimation for operational purchases.",
        "Work authorisation: Nigerian citizen, based in Lagos and available for full-time office or hybrid roles.",
    ]:
        story.append(Paragraph("•  " + item, bullet))

    add_sec("Interests")
    story.append(
        Paragraph(
            "Financial magazines and educative articles; philosophy and motivational reading; "
            "novels; writing poetry. Continued personal study in risk, insurance mathematics, "
            "and practical financial analysis.",
            body,
        )
    )

    add_sec("References")
    story.append(Paragraph("Available on request.", body))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setStrokeColor(RULE)
        canvas.setLineWidth(0.6)
        canvas.line(1.55 * cm, 1.05 * cm, A4[0] - 1.55 * cm, 1.05 * cm)
        canvas.setFont("Times-Italic", 8)
        canvas.setFillColor(LIGHT)
        canvas.drawString(1.55 * cm, 0.7 * cm, "Kelvin Oladotun Esanoluwa  ·  Curriculum Vitae")
        canvas.drawRightString(A4[0] - 1.55 * cm, 0.7 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    docx_path = OUT_DIR / "Kelvin_Oladotun_Esanoluwa_CV.docx"
    pdf_path = OUT_DIR / "Kelvin_Oladotun_Esanoluwa_CV.pdf"
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"Wrote {docx_path}")
    print(f"Wrote {pdf_path}")
